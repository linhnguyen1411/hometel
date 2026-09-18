import io
import uuid
import pytest
import openpyxl
from httpx import AsyncClient, ASGITransport
from app.main import app


# Register a temporary crash route on the app to verify 500 global exception handler returns uniform JSON
@app.get("/api/test-unhandled-crash")
async def crash_endpoint():
    raise RuntimeError("Intentional test crash for resilience verification")


@pytest.mark.asyncio
async def test_global_exception_handling_returns_json():
    """Verify that 404, 422, and unhandled 500 errors all return uniform JSON, never plaintext."""
    # Note: raise_app_exceptions=False allows testing HTTP 500 translation by Starlette/FastAPI middleware
    async with AsyncClient(transport=ASGITransport(app=app, raise_app_exceptions=False), base_url="http://test") as ac:
        # 1. Test 404 Not Found returns uniform JSON
        res_404 = await ac.get("/api/v1/non-existent-endpoint-xyz")
        assert res_404.status_code == 404
        assert "application/json" in res_404.headers.get("content-type", "")
        body_404 = res_404.json()
        assert body_404["success"] is False
        assert "error" in body_404
        assert body_404["error"]["code"] == "HTTP_404"

        # 2. Test 422 Validation Error returns uniform JSON
        res_422 = await ac.post("/api/v1/auth/login", json={"invalid_payload": 123})
        assert res_422.status_code == 422
        assert "application/json" in res_422.headers.get("content-type", "")
        body_422 = res_422.json()
        assert body_422["success"] is False
        assert body_422["error"]["code"] == "VALIDATION_ERROR"

        # 3. Test unhandled 500 error returns uniform JSON (no Starlette plaintext crash)
        res_500 = await ac.get("/api/test-unhandled-crash")
        assert res_500.status_code == 500
        assert "application/json" in res_500.headers.get("content-type", "")
        body_500 = res_500.json()
        assert body_500["success"] is False
        assert body_500["error"]["code"] == "INTERNAL_ERROR"
        assert "Lỗi máy chủ nội bộ. Vui lòng thử lại sau." in body_500["error"]["message"]


@pytest.mark.asyncio
async def test_building_excel_template_and_bulk_import():
    """Verify downloading building Excel template and importing buildings with atomic rollback."""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        login_res = await ac.post("/api/v1/auth/login", json={"email": "owner@homtel.vn", "password": "Owner@123"})
        assert login_res.status_code == 200
        token = login_res.json()["data"]["accessToken"]
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Download template
        tpl_res = await ac.get("/api/v1/buildings/import-template")
        assert tpl_res.status_code == 200
        assert "spreadsheetml.sheet" in tpl_res.headers.get("content-type", "")
        assert len(tpl_res.content) > 1000

        # 2. Test atomic rollback on invalid row in Excel
        bad_bld_name = f"Bad Building {uuid.uuid4().hex[:6]}"
        bad_wb = openpyxl.Workbook()
        bad_ws = bad_wb.active
        bad_ws.append(["Tên tòa nhà", "Địa chỉ", "Quận", "TP", "Tầng", "Phòng", "Loại", "Diện tích", "Giá thuê", "Nội thất", "Sức chứa"])
        bad_ws.append([bad_bld_name, "123 Error St", "Q1", "HCM", 1, "101", "STUDIO", 30, 5000000, "FULLY_FURNISHED", 2])
        bad_ws.append([bad_bld_name, "123 Error St", "Q1", "HCM", 1, "102", "STUDIO", -10, 5000000, "FULLY_FURNISHED", 2])  # Invalid area!
        bad_buf = io.BytesIO()
        bad_wb.save(bad_buf)
        bad_buf.seek(0)

        bad_import_res = await ac.post(
            "/api/v1/buildings/import",
            files={"file": ("bad_buildings.xlsx", bad_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers
        )
        assert bad_import_res.status_code == 422
        bad_data = bad_import_res.json()
        assert bad_data["success"] is False
        assert "3" in bad_data["error"]["message"]

        # 3. Test successful bulk import with unique building name
        good_bld_name = f"Homtel Riverside {uuid.uuid4().hex[:6]}"
        good_wb = openpyxl.Workbook()
        good_ws = good_wb.active
        good_ws.append(["Tên tòa nhà", "Địa chỉ", "Quận", "TP", "Tầng", "Phòng", "Loại", "Diện tích", "Giá thuê", "Nội thất", "Sức chứa"])
        good_ws.append([good_bld_name, "888 Riverside Blvd", "Bình Thạnh", "Hồ Chí Minh", 1, "R101", "STUDIO", 35.5, 8000000, "FULLY_FURNISHED", 2])
        good_ws.append([good_bld_name, "888 Riverside Blvd", "Bình Thạnh", "Hồ Chí Minh", 1, "R102", "ONE_BEDROOM", 48.0, 10500000, "FULLY_FURNISHED", 3])
        good_buf = io.BytesIO()
        good_wb.save(good_buf)
        good_buf.seek(0)

        good_import_res = await ac.post(
            "/api/v1/buildings/import",
            files={"file": ("good_buildings.xlsx", good_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers
        )
        assert good_import_res.status_code in (200, 201), f"Import error: {good_import_res.text}"
        good_data = good_import_res.json()["data"]
        assert good_data["buildingsCreated"] >= 1
        assert good_data["roomsCreated"] >= 2


@pytest.mark.asyncio
async def test_meter_reading_template_and_bulk_import():
    """Verify meter template download, manual reading commit, and bulk Excel import."""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        login_res = await ac.post("/api/v1/auth/login", json={"email": "owner@homtel.vn", "password": "Owner@123"})
        token = login_res.json()["data"]["accessToken"]
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Download meter template
        tpl_res = await ac.get("/api/v1/billing/meters/import-template")
        assert tpl_res.status_code == 200
        assert "spreadsheetml.sheet" in tpl_res.headers.get("content-type", "")

        # 2. Get a room to test meter fetching and manual reading commit
        rooms_res = await ac.get("/api/v1/rooms")
        assert rooms_res.status_code == 200
        rooms = rooms_res.json()["data"]
        assert len(rooms) > 0
        first_room = rooms[0]

        meters_res = await ac.get(f"/api/v1/billing/meters/room/{first_room['id']}", headers=headers)
        assert meters_res.status_code == 200
        meters = meters_res.json()["data"]
        assert len(meters) > 0
        target_meter = meters[0]

        new_val = (target_meter.get("currentReading") or 100) + 15.5
        manual_commit = await ac.post(f"/api/v1/billing/meters/{target_meter['id']}/commit-ocr", json={
            "readingValue": new_val,
            "ocrConfidence": None,  # None means manual reading entered directly
            "rawOcrText": "Ghi chỉ số thủ công bởi quản lý",
            "notes": "Kiểm tra định kỳ trực tiếp"
        }, headers=headers)
        assert manual_commit.status_code == 200
        commit_data = manual_commit.json()["data"]
        assert commit_data["currentReading"] == new_val
        assert commit_data["meterId"] == target_meter["id"]

        # 3. Test bulk meter reading import via Excel
        meter_wb = openpyxl.Workbook()
        meter_ws = meter_wb.active
        meter_ws.append(["Tòa nhà (*)", "Số phòng (*)", "Loại công tơ (DIEN/NUOC) (*)", "Chỉ số mới (*)", "Ngày ghi (YYYY-MM-DD)", "Ghi chú"])
        # Building name from seeded building
        bld_res = await ac.get("/api/v1/buildings")
        bld_name = bld_res.json()["data"][0]["name"]
        meter_type_vn = "DIEN" if target_meter["type"] == "ELECTRICITY" else "NUOC"
        meter_ws.append([bld_name, first_room["roomNumber"], meter_type_vn, new_val + 5.0, "2026-09-17", "Bulk import automated test"])
        meter_buf = io.BytesIO()
        meter_wb.save(meter_buf)
        meter_buf.seek(0)

        bulk_res = await ac.post(
            "/api/v1/billing/meters/bulk-import",
            files={"file": ("bulk_meters.xlsx", meter_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers
        )
        assert bulk_res.status_code == 200, f"Bulk import meters failed: {bulk_res.text}"
        bulk_data = bulk_res.json()["data"]
        assert bulk_data["importedCount"] >= 1
