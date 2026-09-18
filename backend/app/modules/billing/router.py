import io
import uuid
import re
import secrets
import logging
from datetime import datetime, timedelta, date, timezone
from decimal import Decimal
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, status, Query, Header, Request, UploadFile, File, Response
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import select, and_, desc, func
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession
from pydantic import BaseModel

from app.core.config import settings
from app.core.database import get_db
from app.core.security import get_current_user, require_role, TokenPayload
from app.core.authz import (
    get_user_company_ids,
    assert_invoice_access,
    assert_room_access,
    assert_building_access,
)
from app.common.response import success_response, error_response
from .models import Meter, MeterReading, Invoice, InvoiceItem, Payment
from .schemas import OcrScanRequest, CommitOcrReadingRequest, GenerateInvoiceRequest, RecordPaymentRequest, VietQrWebhookPayload
from app.modules.rentals.models import RentalContract
from app.modules.properties.models import Room, Building

logger = logging.getLogger("homtel.billing.webhook")

router = APIRouter(tags=["Billing & Metering"])


class GenerateMonthlyInvoicesPayload(BaseModel):
    buildingId: Optional[str] = None
    period: Optional[str] = None
    billing_month: Optional[str] = None
    dueDate: Optional[str] = None


# =========================================================================
# METERS & OCR SMART VISION SCANNING
# =========================================================================

@router.post("/meters/ocr-scan", summary="Quét ảnh công tơ bằng AI OCR Vision")
async def scan_meter_ocr(
    payload: OcrScanRequest,
    current_user: TokenPayload = Depends(require_role("OWNER", "STAFF", "SUPER_ADMIN")),
    db: AsyncSession = Depends(get_db),
):
    stmt = select(Meter).where(Meter.id == payload.meterId)
    res = await db.execute(stmt)
    meter = res.scalar_one_or_none()
    if not meter:
        return error_response(code="METER_NOT_FOUND", message="Không tìm thấy công tơ", status_code=404)

    await assert_room_access(db, current_user, meter.room_id)

    # Extract reading from base64 string or mock realistic value
    mock_delta = 18.5 if meter.type == "ELECTRICITY" else 4.2
    raw_match = re.search(r"reading=([0-9.]+)", payload.imageBase64OrUrl)
    previous_val = float(meter.current_reading)
    extracted_val = float(raw_match.group(1)) if raw_match else round(previous_val + mock_delta, 1)

    consumption = round(extracted_val - previous_val, 1)

    if extracted_val < previous_val:
        return error_response(
            code="INVALID_READING",
            message=f"Chỉ số mới ({extracted_val}) không thể nhỏ hơn chỉ số cũ ({previous_val})",
            status_code=400
        )

    # Anomaly detection: flag if consumption exceeds normal threshold by 300%
    normal_avg = 50.0 if meter.type == "ELECTRICITY" else 15.0
    is_anomaly = consumption > (normal_avg * 3.0)

    confidence = 0.96 if not is_anomaly else 0.88
    warning = "CẢNH BÁO: Mức tiêu thụ tăng vọt bất thường (>300%) so với bình quân!" if is_anomaly else None

    return success_response({
        "meterId": meter.id,
        "meterType": meter.type,
        "serialNumber": meter.serial_number,
        "previousReading": previous_val,
        "detectedReading": extracted_val,
        "consumption": consumption,
        "confidence": confidence,
        "isAnomaly": is_anomaly,
        "warning": warning,
        "rawOcrText": f"MOD: 2026-IoT-DIGITAL\nTOTAL: {extracted_val} {'kWh' if meter.type == 'ELECTRICITY' else 'm3'}\nCONF: {confidence * 100}%"
    })


@router.get("/meters/import-template", summary="Tải template Excel nhập chỉ số công tơ")
async def download_meter_import_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ChiSoCongTo"
    headers = [
        "Tòa nhà (*)",
        "Số phòng (*)",
        "Loại công tơ (DIEN/NUOC) (*)",
        "Chỉ số mới (*)",
        "Ngày ghi (YYYY-MM-DD)",
        "Ghi chú"
    ]
    ws.append(headers)
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sample_rows = [
        ["Homtel Central Park", "101", "DIEN", 250.5, datetime.now(timezone.utc).strftime("%Y-%m-%d"), "Chốt định kỳ"],
        ["Homtel Central Park", "101", "NUOC", 15.0, datetime.now(timezone.utc).strftime("%Y-%m-%d"), "Chốt định kỳ"],
        ["Homtel Central Park", "102", "DIEN", 310.0, datetime.now(timezone.utc).strftime("%Y-%m-%d"), "Chốt định kỳ"],
        ["Homtel Central Park", "102", "NUOC", 22.0, datetime.now(timezone.utc).strftime("%Y-%m-%d"), "Chốt định kỳ"],
    ]
    for r in sample_rows:
        ws.append(r)

    col_widths = [25, 14, 28, 16, 22, 25]
    for idx, width in enumerate(col_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=homtel_meter_import_template.xlsx"}
    )


@router.post("/meters/bulk-import", summary="Nhập hàng loạt chỉ số công tơ từ Excel")
async def bulk_import_meter_readings(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "STAFF", "SUPER_ADMIN")),
):
    if not file.filename.endswith((".xlsx", ".xls")):
        return error_response(code="INVALID_FILE_TYPE", message="Vui lòng tải lên file Excel (.xlsx)", status_code=400)

    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(filename=io.BytesIO(contents), data_only=True)
        ws = wb.active
    except Exception as e:
        return error_response(code="CORRUPTED_FILE", message=f"Không thể đọc file Excel: {str(e)}", status_code=400)

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return error_response(code="EMPTY_FILE", message="File Excel không có dòng dữ liệu nào", status_code=400)

    # Validate all rows first
    parsed_items = []
    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        bld_name = str(row[0]).strip() if row[0] is not None else ""
        room_num = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        raw_type = str(row[2]).strip().upper() if len(row) > 2 and row[2] is not None else ""
        reading_val = row[3] if len(row) > 3 else None
        reading_date = str(row[4]).strip() if len(row) > 4 and row[4] is not None else datetime.now(timezone.utc).strftime("%Y-%m-%d")
        notes = str(row[5]).strip() if len(row) > 5 and row[5] is not None else "Bulk import Excel"

        if not bld_name:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Thiếu tên tòa nhà", status_code=422)
        if not room_num:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Thiếu số phòng", status_code=422)
        if not raw_type:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Thiếu loại công tơ", status_code=422)

        m_type = "ELECTRICITY" if ("DIEN" in raw_type or "ELEC" in raw_type) else ("WATER" if ("NUOC" in raw_type or "WAT" in raw_type) else None)
        if not m_type:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Loại công tơ '{raw_type}' không hợp lệ (phải là DIEN hoặc NUOC)", status_code=422)

        if reading_val is None:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Thiếu chỉ số mới", status_code=422)
        try:
            val = float(reading_val)
            if val < 0:
                raise ValueError()
        except (ValueError, TypeError):
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Chỉ số '{reading_val}' không hợp lệ", status_code=422)

        parsed_items.append({
            "row_idx": row_idx,
            "bld_name": bld_name,
            "room_num": room_num,
            "meter_type": m_type,
            "reading_val": val,
            "reading_date": reading_date,
            "notes": notes,
        })

    # Validate database records and perform atomic updates
    updated_count = 0
    user_comp_ids = get_user_company_ids(current_user)
    try:
        for item in parsed_items:
            r_stmt = (
                select(Room)
                .join(Building, Room.building_id == Building.id)
                .where(
                    func.lower(Building.name) == item["bld_name"].lower(),
                    Room.room_number == item["room_num"]
                )
            )
            if current_user.role != "SUPER_ADMIN":
                r_stmt = r_stmt.where(Building.company_id.in_(user_comp_ids))
            r_res = await db.execute(r_stmt)
            room = r_res.scalar_one_or_none()
            if not room:
                raise ValueError(f"Dòng {item['row_idx']}: Không tìm thấy phòng '{item['room_num']}' trong tòa '{item['bld_name']}' hoặc bạn không có quyền truy cập")

            m_stmt = select(Meter).where(Meter.room_id == room.id, Meter.type == item["meter_type"])
            m_res = await db.execute(m_stmt)
            meter = m_res.scalar_one_or_none()
            if not meter:
                meter = Meter(
                    id=f"mtr_{uuid.uuid4().hex[:12]}",
                    room_id=room.id,
                    type=item["meter_type"],
                    serial_number=f"{'EM' if item['meter_type'] == 'ELECTRICITY' else 'WM'}-{room.room_number}-{secrets.randbelow(9000)+1000}",
                    initial_reading=Decimal(str(item["reading_val"])),
                    current_reading=Decimal(str(item["reading_val"])),
                    status="ACTIVE"
                )
                db.add(meter)
                await db.flush()

            prev_val = float(meter.current_reading)
            new_val = float(item["reading_val"])
            if new_val < prev_val:
                raise ValueError(f"Dòng {item['row_idx']}: Chỉ số mới ({new_val}) nhỏ hơn chỉ số hiện tại ({prev_val}) của phòng {room.room_number}")

            consumption = round(new_val - prev_val, 1)
            reading = MeterReading(
                id=f"rdg_{uuid.uuid4().hex[:12]}",
                meter_id=meter.id,
                previous_reading=prev_val,
                reading_value=new_val,
                consumption=consumption,
                reading_date=item["reading_date"],
                recorded_by=current_user.userId,
                notes=item["notes"],
                ocr_confidence=None
            )
            db.add(reading)
            meter.current_reading = Decimal(str(new_val))
            updated_count += 1

        await db.commit()
    except Exception as e:
        await db.rollback()
        return error_response(code="IMPORT_METER_ERROR", message=f"Lỗi khi import: {str(e)}", status_code=400)

    return success_response({
        "message": f"Đã cập nhật thành công {updated_count} chỉ số công tơ!",
        "importedCount": updated_count,
    })


@router.post("/meters/{meter_id}/commit-ocr", summary="Lưu chốt số công tơ & tùy chọn tạo hóa đơn tự động")
async def commit_ocr_reading(
    meter_id: str,
    payload: CommitOcrReadingRequest,
    current_user: TokenPayload = Depends(require_role("OWNER", "STAFF", "SUPER_ADMIN")),
    db: AsyncSession = Depends(get_db),
):
    stmt = select(Meter).where(Meter.id == meter_id)
    res = await db.execute(stmt)
    meter = res.scalar_one_or_none()
    if not meter:
        return error_response(code="METER_NOT_FOUND", message="Không tìm thấy công tơ", status_code=404)

    await assert_room_access(db, current_user, meter.room_id)

    previous_val = float(meter.current_reading)
    new_val = float(payload.readingValue)
    if new_val < previous_val:
        return error_response(code="INVALID_READING", message="Chỉ số mới không thể nhỏ hơn chỉ số cũ", status_code=400)

    consumption = round(new_val - previous_val, 1)
    now_iso = datetime.now(timezone.utc).isoformat()

    reading = MeterReading(
        id=f"rdg_{uuid.uuid4().hex[:12]}",
        meter_id=meter.id,
        previous_reading=previous_val,
        reading_value=new_val,
        consumption=consumption,
        reading_date=payload.readingDate or datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        recorded_by=current_user.userId,
        notes=payload.notes or ("Ghi số thủ công" if payload.ocrConfidence is None else None),
        image_url=payload.imageUrl,
        ocr_confidence=payload.ocrConfidence,
        ocr_raw_text=payload.ocrRawText
    )
    db.add(reading)
    meter.current_reading = Decimal(str(new_val))

    invoice_data = None
    if payload.autoDraftInvoice:
        # Find active contract for room
        c_stmt = select(RentalContract).where(
            and_(
                RentalContract.room_id == meter.room_id,
                RentalContract.status == "ACTIVE"
            )
        )
        c_res = await db.execute(c_stmt)
        contract = c_res.scalar_one_or_none()
        if contract:
            rate = Decimal("4000.00") if meter.type == "ELECTRICITY" else Decimal("18000.00")
            utility_amount = Decimal(str(round(consumption * float(rate), 2)))
            contract_rent = Decimal(str(contract.rent_amount))
            total_bill = contract_rent + utility_amount
            inv_id = f"inv_{uuid.uuid4().hex[:12]}"
            inv_num = f"INV-{datetime.now(timezone.utc).strftime('%Y%m')}-{secrets.randbelow(9000) + 1000}"

            new_invoice = Invoice(
                id=inv_id,
                invoice_number=inv_num,
                tenant_id=contract.tenant_id,
                contract_id=contract.id,
                company_id=contract.company_id,
                room_id=contract.room_id,
                issue_date=datetime.now(timezone.utc).strftime("%Y-%m-%d"),
                due_date=(datetime.now(timezone.utc) + timedelta(days=7)).strftime("%Y-%m-%d"),
                billing_month=datetime.now(timezone.utc).strftime("%Y-%m"),
                subtotal=total_bill,
                total=total_bill,
                paid_amount=Decimal("0.00"),
                outstanding_amount=total_bill,
                status="DRAFT",
                notes="Hóa đơn tháng tự động chốt từ ảnh đồng hồ công tơ OCR."
            )
            db.add(new_invoice)

            # Add Rent Item
            db.add(InvoiceItem(
                id=f"item_{uuid.uuid4().hex[:12]}",
                invoice_id=inv_id,
                type="RENT",
                description=f"Tiền thuê căn hộ tháng {datetime.now(timezone.utc).strftime('%m/%Y')}",
                quantity=Decimal("1.0"),
                unit_price=contract_rent,
                amount=contract_rent
            ))

            # Add Utility Item
            db.add(InvoiceItem(
                id=f"item_{uuid.uuid4().hex[:12]}",
                invoice_id=inv_id,
                type="ELECTRICITY" if meter.type == "ELECTRICITY" else "WATER",
                description=f"Tiền {'điện' if meter.type == 'ELECTRICITY' else 'nước'} ({previous_val} → {new_val} = {consumption} {'kWh' if meter.type == 'ELECTRICITY' else 'm3'})",
                quantity=Decimal(str(consumption)),
                unit_price=rate,
                amount=utility_amount
            ))

            invoice_data = {
                "id": inv_id,
                "invoiceNumber": inv_num,
                "total": float(total_bill),
                "status": "DRAFT"
            }

    await db.commit()

    return success_response({
        "meterId": meter.id,
        "currentReading": float(meter.current_reading),
        "consumption": consumption,
        "readingId": reading.id,
        "autoDraftedInvoice": invoice_data
    })


@router.get("/meters/room/{room_id}", summary="Lấy danh sách công tơ theo phòng")
async def get_room_meters(
    room_id: str,
    current_user: TokenPayload = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await assert_room_access(db, current_user, room_id)
    stmt = select(Meter).where(Meter.room_id == room_id).options(selectinload(Meter.readings))
    res = await db.execute(stmt)
    meters = res.scalars().all()

    data = [
        {
            "id": m.id,
            "roomId": m.room_id,
            "type": m.type,
            "serialNumber": m.serial_number,
            "currentReading": float(m.current_reading) if m.current_reading is not None else 0.0,
            "status": m.status,
            "recentReadings": [
                {
                    "id": r.id,
                    "previousReading": float(r.previous_reading) if r.previous_reading is not None else 0.0,
                    "readingValue": float(r.reading_value) if r.reading_value is not None else 0.0,
                    "consumption": float(r.consumption) if r.consumption is not None else 0.0,
                    "readingDate": r.reading_date,
                    "ocrConfidence": r.ocr_confidence
                }
                for r in m.readings[:5]
            ]
        }
        for m in meters
    ]
    return success_response(data)


# =========================================================================
# INVOICES, VIETQR & PAYMENTS
# =========================================================================

@router.get("/invoices", summary="Danh sách hóa đơn")
async def list_invoices(
    status: Optional[str] = Query(None),
    current_user: TokenPayload = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    stmt = select(Invoice).options(selectinload(Invoice.items), selectinload(Invoice.payments))

    if current_user.role == "TENANT":
        stmt = stmt.where(Invoice.tenant_id == current_user.userId)
    elif current_user.role in ("OWNER", "STAFF") and current_user.memberships:
        company_ids = [m.companyId for m in current_user.memberships]
        stmt = stmt.where(Invoice.company_id.in_(company_ids))

    if status:
        stmt = stmt.where(Invoice.status == status)

    stmt = stmt.order_by(desc(Invoice.created_at))
    res = await db.execute(stmt)
    invoices = res.scalars().all()

    data = [
        {
            "id": inv.id,
            "invoice_number": inv.invoice_number,
            "tenant_id": inv.tenant_id,
            "contract_id": inv.contract_id,
            "room_id": inv.room_id,
            "issue_date": inv.issue_date,
            "due_date": inv.due_date,
            "billing_month": inv.billing_month,
            "subtotal": inv.subtotal,
            "total": inv.total,
            "paid_amount": inv.paid_amount,
            "outstanding_amount": inv.outstanding_amount,
            "status": inv.status,
            "items": [
                {
                    "id": item.id,
                    "type": item.type,
                    "description": item.description,
                    "quantity": item.quantity,
                    "unitPrice": item.unit_price,
                    "amount": item.amount
                }
                for item in inv.items
            ]
        }
        for inv in invoices
    ]
    return success_response(data)


@router.post("/invoices/generate-monthly", summary="Tạo hóa đơn định kỳ hàng tháng cho toàn bộ hợp đồng active của tòa nhà")
async def generate_monthly_invoices(
    payload: GenerateMonthlyInvoicesPayload,
    current_user: TokenPayload = Depends(require_role("OWNER", "STAFF", "SUPER_ADMIN")),
    db: AsyncSession = Depends(get_db),
):
    period = payload.billing_month or payload.period or datetime.now(timezone.utc).strftime("%Y-%m")
    
    # Query active contracts
    contract_stmt = (
        select(RentalContract)
        .join(Room, RentalContract.room_id == Room.id)
        .options(selectinload(RentalContract.room))
        .where(RentalContract.status == "ACTIVE")
    )
    if current_user.role != "SUPER_ADMIN":
        user_companies = get_user_company_ids(current_user)
        if payload.buildingId:
            await assert_building_access(db, current_user, payload.buildingId)
            contract_stmt = contract_stmt.where(Room.building_id == payload.buildingId)
        else:
            contract_stmt = contract_stmt.where(RentalContract.company_id.in_(user_companies))
    elif payload.buildingId:
        contract_stmt = contract_stmt.where(Room.building_id == payload.buildingId)

    contract_res = await db.execute(contract_stmt)
    contracts = contract_res.scalars().all()

    created_invoices = []
    skipped_count = 0
    created_count = 0

    from sqlalchemy.exc import IntegrityError

    for contract in contracts:
        try:
            async with db.begin_nested():
                # Idempotency check: has an invoice already been generated for this contract & period?
                inv_check = await db.execute(
                    select(Invoice).where(
                        Invoice.contract_id == contract.id,
                        Invoice.billing_month == period
                    ).with_for_update()
                )
                if inv_check.scalar_one_or_none():
                    skipped_count += 1
                    continue

                clean_period = period.replace("-", "")
                room_num = contract.room.room_number if contract.room else "ROOM"
                inv_num = f"INV-{clean_period}-{room_num}-{secrets.randbelow(9000)+1000}"
                
                due_date = payload.dueDate or f"{period}-10"
                rent_amount = Decimal(str(contract.rent_amount))

                new_inv = Invoice(
                    id=f"inv_{uuid.uuid4().hex[:12]}",
                    invoice_number=inv_num,
                    tenant_id=contract.tenant_id,
                    contract_id=contract.id,
                    company_id=contract.company_id,
                    room_id=contract.room_id,
                    issue_date=date.today().isoformat(),
                    due_date=due_date,
                    billing_month=period,
                    subtotal=rent_amount,
                    discount=Decimal("0.00"),
                    tax=Decimal("0.00"),
                    total=rent_amount,
                    paid_amount=Decimal("0.00"),
                    outstanding_amount=rent_amount,
                    status="DRAFT",
                    notes=f"Hóa đơn tiền phòng định kỳ tháng {period}"
                )
                db.add(new_inv)
                await db.flush()

                item = InvoiceItem(
                    id=f"item_{uuid.uuid4().hex[:12]}",
                    invoice_id=new_inv.id,
                    type="RENT",
                    description=f"Tiền thuê căn hộ {room_num} (Tháng {period})",
                    quantity=Decimal("1.0"),
                    unit_price=rent_amount,
                    amount=rent_amount,
                )
                db.add(item)
                created_invoices.append(new_inv)
                created_count += 1
        except IntegrityError:
            # Another concurrent transaction inserted the invoice for this period
            skipped_count += 1
            continue

    await db.commit()

    return success_response({
        "buildingId": payload.buildingId,
        "period": period,
        "totalActiveContracts": len(contracts),
        "createdCount": created_count,
        "skippedCount": skipped_count,
        "invoices_created": created_count,
        "invoices_skipped": skipped_count,
        "invoices": [
            {
                "id": i.id,
                "invoiceNumber": i.invoice_number,
                "tenantId": i.tenant_id,
                "total": i.total,
                "status": i.status
            }
            for i in created_invoices
        ]
    }, status_code=status.HTTP_201_CREATED if created_count > 0 else status.HTTP_200_OK)


@router.get("/invoices/{invoice_id}/vietqr", summary="Lấy mã thanh toán VietQR NAPAS 247")
async def get_invoice_vietqr(
    invoice_id: str,
    current_user: TokenPayload = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    if not settings.PAYMENT_GATEWAY_ENABLED:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Tính năng đang tạm ngưng, sẽ kích hoạt sau khi hoàn tất đăng ký doanh nghiệp",
        )

    inv = await assert_invoice_access(db, current_user, invoice_id)

    bank_id = "MB"  # Military Commercial Joint Stock Bank
    account_no = "0905111001"
    account_name = "HOMTEL DA NANG MANAGEMENT"
    amount = int(inv.outstanding_amount if inv.outstanding_amount > 0 else inv.total)
    description = f"HOMTEL {inv.invoice_number}"

    # Generate standard NAPAS 247 Quicklink URL
    vietqr_url = f"https://img.vietqr.io/image/{bank_id}-{account_no}-compact2.png?amount={amount}&addInfo={description}&accountName={account_name}"

    return success_response({
        "invoiceId": inv.id,
        "invoiceNumber": inv.invoice_number,
        "amount": amount,
        "bankId": bank_id,
        "bankName": "MB Bank (Ngân hàng TMCP Quân Đội)",
        "accountNumber": account_no,
        "accountName": account_name,
        "transferContent": description,
        "qrImageUrl": vietqr_url,
        "status": inv.status
    })


@router.post("/invoices/webhook/vietqr", summary="Webhook nhận thông báo biến động số dư ngân hàng VietQR")
async def vietqr_webhook(
    payload: VietQrWebhookPayload,
    request: Request,
    db: AsyncSession = Depends(get_db),
    x_webhook_secret: Optional[str] = Header(None, alias="X-Webhook-Secret"),
):
    """
    Webhook xử lý biến động số dư từ cổng thanh toán VietQR (SePay/Casso/PayOS/Ngân hàng).
    Bắt buộc xác thực Secret key qua header X-Webhook-Secret để chống giả mạo giao dịch tài chính.
    """
    if not settings.PAYMENT_GATEWAY_ENABLED:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Tính năng đang tạm ngưng, sẽ kích hoạt sau khi hoàn tất đăng ký doanh nghiệp",
        )
    client_ip = request.client.host if request.client else "unknown"

    if not x_webhook_secret or not secrets.compare_digest(x_webhook_secret, settings.VIETQR_WEBHOOK_SECRET):
        logger.warning(
            f"[SECURITY ALERT] Unauthorized VietQR webhook attempt from IP: {client_ip}. "
            f"Provided Secret: {'PRESENT (MISMATCH)' if x_webhook_secret else 'MISSING'}"
        )
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="INVALID_WEBHOOK_SECRET: Chữ ký hoặc secret webhook không hợp lệ.",
        )

    # Match invoice from transfer description (e.g. "HOMTEL INV-202609-1234" or "HOMTEL INV-202609-101-A-1234")
    inv_match = re.search(r"INV-[A-Za-z0-9]+(?:-[A-Za-z0-9]+)*", payload.description)
    if not inv_match:
        return success_response({
            "status": "SKIPPED",
            "message": "Không tìm thấy mã hóa đơn hợp lệ trong nội dung chuyển khoản."
        })

    invoice_number = inv_match.group(0)
    stmt = select(Invoice).where(Invoice.invoice_number == invoice_number).with_for_update()
    res = await db.execute(stmt)
    inv = res.scalar_one_or_none()
    if not inv:
        return success_response({
            "status": "NOT_FOUND",
            "message": f"Hóa đơn {invoice_number} không tồn tại."
        })

    pay_amt = Decimal(str(payload.amount))
    if pay_amt <= Decimal("0"):
        return success_response({
            "status": "SKIPPED",
            "message": "Số tiền thanh toán phải lớn hơn 0."
        })

    # Check duplicate transaction reference
    tx_ref = payload.transactionId or payload.referenceNumber
    if tx_ref:
        p_stmt = select(Payment).where(Payment.transaction_reference == tx_ref).with_for_update()
        p_res = await db.execute(p_stmt)
        if p_res.scalar_one_or_none():
            return success_response({
                "status": "DUPLICATE_SKIPPED",
                "message": "Giao dịch này đã được ghi nhận trước đó."
            })

    # Record payment
    payment = Payment(
        id=f"pay_{uuid.uuid4().hex[:12]}",
        invoice_id=inv.id,
        tenant_id=inv.tenant_id,
        company_id=inv.company_id,
        amount=pay_amt,
        method="ONLINE",
        status="SUCCESS",
        transaction_reference=tx_ref or f"VQR-{secrets.randbelow(900000) + 100000}",
        paid_at=datetime.now(timezone.utc).isoformat(),
        notes=f"Tự động gạch nợ VietQR: {payload.description}"
    )
    db.add(payment)

    # Update invoice paid and outstanding amounts
    inv.paid_amount = Decimal(str(inv.paid_amount)) + pay_amt
    inv.outstanding_amount = max(Decimal("0.00"), Decimal(str(inv.total)) - inv.paid_amount)
    if inv.outstanding_amount <= Decimal("0.00"):
        inv.status = "PAID"
    else:
        inv.status = "PARTIALLY_PAID"

    await db.commit()

    return success_response({
        "status": "RECONCILED",
        "invoiceNumber": inv.invoice_number,
        "paidAmount": float(inv.paid_amount),
        "outstandingAmount": float(inv.outstanding_amount),
        "invoiceStatus": inv.status
    })


@router.post("/payments", summary="Tạo thanh toán cho hóa đơn", status_code=201)
async def record_payment(
    payload: RecordPaymentRequest,
    current_user: TokenPayload = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    inv = await assert_invoice_access(db, current_user, payload.invoiceId, for_update=True)

    pay_amt = Decimal(str(payload.amount))
    if pay_amt <= Decimal("0"):
        return error_response(code="INVALID_AMOUNT", message="Số tiền thanh toán phải lớn hơn 0", status_code=400)

    outstanding = Decimal(str(inv.outstanding_amount))
    if pay_amt > outstanding:
        return error_response(
            code="OVERPAYMENT_NOT_ALLOWED",
            message=f"Số tiền thanh toán ({pay_amt}) vượt quá số tiền còn nợ ({outstanding})",
            status_code=400
        )

    payment = Payment(
        id=f"pay_{uuid.uuid4().hex[:12]}",
        invoice_id=inv.id,
        tenant_id=inv.tenant_id,
        company_id=inv.company_id,
        amount=pay_amt,
        method=payload.method,
        status="SUCCESS",
        transaction_reference=payload.transactionReference or f"PAY-{secrets.randbelow(900000) + 100000}",
        paid_at=datetime.now(timezone.utc).isoformat(),
        notes=payload.notes
    )
    db.add(payment)

    inv.paid_amount = Decimal(str(inv.paid_amount)) + pay_amt
    inv.outstanding_amount = max(Decimal("0.00"), Decimal(str(inv.total)) - inv.paid_amount)
    if inv.outstanding_amount <= Decimal("0.00"):
        inv.status = "PAID"
    else:
        inv.status = "PARTIALLY_PAID"

    await db.commit()
    return success_response({
        "id": payment.id,
        "amount": float(payment.amount),
        "invoiceStatus": inv.status,
        "outstandingAmount": float(inv.outstanding_amount)
    }, status_code=201)


billing_router = router
