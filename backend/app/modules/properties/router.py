import io
import re
import secrets
import uuid
from typing import Optional, List
from datetime import date, datetime
from fastapi import APIRouter, Depends, Query, HTTPException, status, UploadFile, File, Response
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, desc, or_
from sqlalchemy.orm import selectinload
from pydantic import BaseModel, ConfigDict, Field

from app.core.database import get_db
from app.core.security import get_optional_user, get_current_user, require_role, TokenPayload
from app.core.authz import (
    resolve_owner_company_id,
    assert_building_access,
    assert_private_building_access,
    assert_public_building_access,
    assert_room_access,
    get_user_company_ids,
)
from app.common.response import success_response, error_response
from app.modules.rentals.models import RentalContract
from app.modules.billing.models import Meter
from .service import PropertyService
from .models import Building, Floor, Room, BuildingConfiguration
from .schemas import CreateBuildingRequest, CreateRoomRequest

router = APIRouter(tags=["Properties & Inventory"])


class CreateBuildingPayload(BaseModel):
    companyId: Optional[str] = None
    name: str
    slug: Optional[str] = None
    description: Optional[str] = None
    address: str
    city: Optional[str] = "Hồ Chí Minh"
    district: Optional[str] = None
    ward: Optional[str] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    imageUrl: Optional[str] = None
    total_floors: Optional[int] = None


class UpdateBuildingPayload(BaseModel):
    name: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    district: Optional[str] = None
    ward: Optional[str] = None
    description: Optional[str] = None
    imageUrl: Optional[str] = None
    status: Optional[str] = None


class CreateFloorPayload(BaseModel):
    floorNumber: int
    name: str
    description: Optional[str] = None


class UpdateFloorPayload(BaseModel):
    floorNumber: Optional[int] = None
    name: Optional[str] = None
    description: Optional[str] = None


class CreateRoomPayload(BaseModel):
    buildingId: Optional[str] = None
    building_id: Optional[str] = None
    floorId: Optional[str] = None
    floor_id: Optional[str] = None
    roomNumber: Optional[str] = None
    room_number: Optional[str] = None
    slug: Optional[str] = None
    roomType: Optional[str] = None
    room_type: Optional[str] = None
    area: Optional[float] = None
    area_sqm: Optional[float] = None
    baseRent: Optional[float] = None
    base_price_monthly: Optional[float] = None
    capacity: Optional[int] = 2
    max_occupants: Optional[int] = None
    status: Optional[str] = "AVAILABLE"
    furnishing: Optional[str] = "FULLY_FURNISHED"
    description: Optional[str] = None
    amenities: Optional[list] = None
    images: Optional[list] = None


class UpdateRoomPayload(BaseModel):
    model_config = ConfigDict(populate_by_name=True)

    roomNumber: Optional[str] = Field(None, alias="room_number")
    room_number: Optional[str] = None
    roomType: Optional[str] = Field(None, alias="room_type")
    room_type: Optional[str] = None
    area: Optional[float] = None
    baseRent: Optional[float] = Field(None, alias="base_rent")
    base_rent: Optional[float] = None
    capacity: Optional[int] = None
    status: Optional[str] = None
    furnishing: Optional[str] = None
    description: Optional[str] = None
    amenities: Optional[list] = None
    images: Optional[list] = None


class CreateBuildingConfigPayload(BaseModel):
    model_config = ConfigDict(populate_by_name=True)

    effectiveFrom: Optional[str] = Field(None, alias="effective_from")
    effective_from: Optional[str] = None
    electricityUnitPrice: Optional[float] = Field(3500, alias="electricity_unit_price")
    electricity_unit_price: Optional[float] = None
    waterUnitPrice: Optional[float] = Field(15000, alias="water_unit_price")
    water_unit_price: Optional[float] = None
    internetPrice: Optional[float] = Field(100000, alias="internet_price")
    internet_price: Optional[float] = None
    garbagePrice: Optional[float] = Field(50000, alias="garbage_price")
    garbage_price: Optional[float] = None
    parkingFeeMotorbike: Optional[float] = Field(100000, alias="parking_fee_motorbike")
    parking_fee_motorbike: Optional[float] = None
    parkingFeeCar: Optional[float] = Field(800000, alias="parking_fee_car")
    parking_fee_car: Optional[float] = None
    cleaningFee: Optional[float] = Field(150000, alias="cleaning_fee")
    cleaning_fee: Optional[float] = None
    notes: Optional[str] = None


# ==========================================
# BUILDINGS
# ==========================================

@router.get("/buildings", summary="Danh sách tòa nhà (Public & Filter)")
async def list_buildings(
    companyId: Optional[str] = Query(None),
    status: Optional[str] = Query("ACTIVE"),
    search: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    pageSize: int = Query(50, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
):
    offset = (page - 1) * pageSize
    buildings = await PropertyService.list_buildings(
        db,
        company_id=companyId,
        status_filter=status,
        search=search,
        limit=pageSize,
        offset=offset,
    )
    return success_response(buildings)


@router.get("/buildings/slug/{slug}", summary="Chi tiết tòa nhà theo slug")
async def get_building_by_slug(slug: str, db: AsyncSession = Depends(get_db)):
    stmt = (
        select(Building)
        .options(selectinload(Building.floors), selectinload(Building.rooms))
        .where(Building.slug == slug.lower())
    )
    res = await db.execute(stmt)
    bld = res.scalar_one_or_none()
    if not bld:
        raise HTTPException(status_code=404, detail="Building not found")

    available_rooms = [
        {
            "id": r.id,
            "roomNumber": r.room_number,
            "slug": r.slug,
            "roomType": r.room_type,
            "area": float(r.area),
            "baseRent": float(r.base_rent),
            "capacity": r.capacity,
            "status": r.status,
            "furnishing": r.furnishing,
            "images": r.images or [],
            "amenities": r.amenities or [],
        }
        for r in bld.rooms
        if r.status == "AVAILABLE"
    ]

    return success_response({
        "id": bld.id,
        "companyId": bld.company_id,
        "name": bld.name,
        "slug": bld.slug,
        "description": bld.description,
        "address": bld.address,
        "city": bld.city,
        "district": bld.district,
        "ward": bld.ward,
        "imageUrl": bld.image_url,
        "status": bld.status,
        "floors": [{"id": f.id, "floorNumber": f.floor_number, "name": f.name} for f in bld.floors],
        "availableRooms": available_rooms,
    })


@router.get("/buildings/import-template", summary="Tải file mẫu Excel import tòa nhà & phòng")
async def download_building_import_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DS_Phong_ToaNha"

    headers = [
        "Tên tòa nhà (*)",
        "Địa chỉ (*)",
        "Quận/Huyện",
        "Tỉnh/Thành phố",
        "Tầng (*)",
        "Số phòng (*)",
        "Loại phòng (*)",
        "Diện tích (m2) (*)",
        "Giá thuê/tháng (VND) (*)",
        "Tình trạng nội thất",
        "Sức chứa (người)"
    ]
    ws.append(headers)

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sample_rows = [
        ["Homtel Central Park", "208 Nguyễn Hữu Cảnh, P.22", "Bình Thạnh", "Hồ Chí Minh", 1, "101", "STUDIO", 32.0, 7500000, "FULLY_FURNISHED", 2],
        ["Homtel Central Park", "208 Nguyễn Hữu Cảnh, P.22", "Bình Thạnh", "Hồ Chí Minh", 1, "102", "ONE_BEDROOM", 45.0, 9500000, "FULLY_FURNISHED", 3],
        ["Homtel Central Park", "208 Nguyễn Hữu Cảnh, P.22", "Bình Thạnh", "Hồ Chí Minh", 2, "201", "TWO_BEDROOM", 65.0, 14000000, "FULLY_FURNISHED", 4],
        ["Homtel Central Park", "208 Nguyễn Hữu Cảnh, P.22", "Bình Thạnh", "Hồ Chí Minh", 2, "202", "STUDIO", 30.0, 7000000, "SEMI_FURNISHED", 2],
    ]
    for row in sample_rows:
        ws.append(row)

    col_widths = [25, 30, 16, 16, 10, 12, 16, 18, 22, 22, 16]
    for idx, width in enumerate(col_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=homtel_building_import_template.xlsx"}
    )


@router.post("/buildings/import", summary="Import hàng loạt tòa nhà, tầng & phòng từ Excel")
async def import_buildings_from_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "SUPER_ADMIN")),
):
    if not file.filename.endswith((".xlsx", ".xls")):
        return error_response(code="INVALID_FILE_TYPE", message="Vui lòng tải lên file Excel (.xlsx)", status_code=400)

    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(filename=io.BytesIO(contents), data_only=True)
        ws = wb.active
    except Exception as e:
        return error_response(code="CORRUPTED_FILE", message=f"Không thể đọc file Excel: {str(e)}", status_code=400)

    comp_id = resolve_owner_company_id(current_user)

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return error_response(code="EMPTY_FILE", message="File Excel không có dòng dữ liệu nào", status_code=400)

    # Validate all rows first
    parsed_rows = []
    seen_room_keys = set()

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue

        bld_name = str(row[0]).strip() if row[0] is not None else ""
        address = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        district = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        city = str(row[3]).strip() if len(row) > 3 and row[3] is not None else "Hồ Chí Minh"
        floor_val = row[4] if len(row) > 4 else None
        room_num = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
        room_type = str(row[6]).strip().upper() if len(row) > 6 and row[6] is not None else "STUDIO"
        area_val = row[7] if len(row) > 7 else None
        rent_val = row[8] if len(row) > 8 else None
        furnishing = str(row[9]).strip().upper() if len(row) > 9 and row[9] is not None else "FULLY_FURNISHED"
        capacity_val = row[10] if len(row) > 10 else 2

        if not bld_name:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Thiếu tên tòa nhà", status_code=422)
        if not address:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Thiếu địa chỉ tòa nhà", status_code=422)
        if floor_val is None:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Thiếu số tầng", status_code=422)
        try:
            floor_number = int(floor_val)
        except (ValueError, TypeError):
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Số tầng '{floor_val}' không hợp lệ (phải là số nguyên)", status_code=422)
        if not room_num:
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Thiếu số phòng", status_code=422)

        try:
            area = float(area_val) if area_val is not None else 25.0
            if area <= 0:
                raise ValueError()
        except (ValueError, TypeError):
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Diện tích phòng '{area_val}' không hợp lệ", status_code=422)

        try:
            base_rent = float(rent_val) if rent_val is not None else 5000000.0
            if base_rent < 0:
                raise ValueError()
        except (ValueError, TypeError):
            return error_response(code="VALIDATION_ERROR", message=f"Dòng {row_idx}: Giá thuê '{rent_val}' không hợp lệ", status_code=422)

        try:
            capacity = int(capacity_val) if capacity_val is not None else 2
        except (ValueError, TypeError):
            capacity = 2

        key = (bld_name.lower(), floor_number, room_num.lower())
        if key in seen_room_keys:
            return error_response(code="DUPLICATE_ROOM", message=f"Dòng {row_idx}: Phòng '{room_num}' tại tầng {floor_number} của tòa '{bld_name}' bị trùng lặp trong file", status_code=422)
        seen_room_keys.add(key)

        parsed_rows.append({
            "row_idx": row_idx,
            "bld_name": bld_name,
            "address": address,
            "district": district,
            "city": city,
            "floor_number": floor_number,
            "room_num": room_num,
            "room_type": room_type if room_type in ["STUDIO", "ONE_BEDROOM", "TWO_BEDROOM", "PENTHOUSE", "DUPLEX"] else "STUDIO",
            "area": area,
            "base_rent": base_rent,
            "furnishing": furnishing if furnishing in ["FULLY_FURNISHED", "SEMI_FURNISHED", "UNFURNISHED"] else "FULLY_FURNISHED",
            "capacity": capacity,
        })

    # All rows validated! Now atomic database execution
    buildings_cache = {}
    floors_cache = {}
    buildings_created = 0
    floors_created = 0
    rooms_created = 0

    try:
        for r in parsed_rows:
            bld_key = r["bld_name"].lower()
            if bld_key not in buildings_cache:
                b_stmt = select(Building).where(
                    Building.company_id == comp_id,
                    func.lower(Building.name) == bld_key
                )
                b_res = await db.execute(b_stmt)
                bld = b_res.scalar_one_or_none()
                if not bld:
                    slug_clean = re.sub(r'[^a-zA-Z0-9]+', '-', r["bld_name"].lower()).strip('-') + f"-{secrets.token_hex(3)}"
                    bld = Building(
                        company_id=comp_id,
                        name=r["bld_name"],
                        slug=slug_clean,
                        address=r["address"],
                        district=r["district"],
                        city=r["city"],
                        image_url="https://images.unsplash.com/photo-1545324418-cc1a3fa10c00?auto=format&fit=crop&w=1000&q=80",
                        status="ACTIVE"
                    )
                    db.add(bld)
                    await db.flush()
                    buildings_created += 1

                    cfg = BuildingConfiguration(
                        building_id=bld.id,
                        version=1,
                        effective_from=date.today(),
                        electricity_unit_price=3500,
                        water_unit_price=15000,
                        internet_price=100000,
                        garbage_price=50000,
                        parking_fee_motorbike=120000,
                        parking_fee_car=1200000,
                        cleaning_fee=50000,
                    )
                    db.add(cfg)
                buildings_cache[bld_key] = bld

            bld = buildings_cache[bld_key]

            # Floor
            flr_key = (bld.id, r["floor_number"])
            if flr_key not in floors_cache:
                f_stmt = select(Floor).where(Floor.building_id == bld.id, Floor.floor_number == r["floor_number"])
                f_res = await db.execute(f_stmt)
                flr = f_res.scalar_one_or_none()
                if not flr:
                    flr = Floor(
                        building_id=bld.id,
                        floor_number=r["floor_number"],
                        name=f"Tầng {r['floor_number']}"
                    )
                    db.add(flr)
                    await db.flush()
                    floors_created += 1
                floors_cache[flr_key] = flr

            flr = floors_cache[flr_key]

            # Check room existence
            rm_stmt = select(Room).where(Room.building_id == bld.id, Room.room_number == r["room_num"])
            rm_res = await db.execute(rm_stmt)
            existing_room = rm_res.scalar_one_or_none()
            if existing_room:
                raise ValueError(f"Dòng {r['row_idx']}: Phòng '{r['room_num']}' đã tồn tại trong tòa nhà '{bld.name}'")

            r_slug = f"{bld.slug}-{r['room_num'].lower()}-{secrets.token_hex(3)}"
            new_room = Room(
                company_id=comp_id,
                building_id=bld.id,
                floor_id=flr.id,
                room_number=r["room_num"],
                slug=r_slug,
                room_type=r["room_type"],
                area=r["area"],
                base_rent=r["base_rent"],
                capacity=r["capacity"],
                status="AVAILABLE",
                furnishing=r["furnishing"],
                description=f"Căn hộ {r['room_num']}, diện tích {r['area']}m2, {r['furnishing'].lower()}.",
                amenities=["Điều hòa", "Nóng lạnh", "Tủ lạnh", "Giường nệm", "Tủ quần áo"],
                images=["https://images.unsplash.com/photo-1522708323590-d24dbb6b0267?auto=format&fit=crop&w=800&q=80"]
            )
            db.add(new_room)
            await db.flush()
            rooms_created += 1

            # Initialize meters
            m_elec = Meter(
                id=f"mtr_{uuid.uuid4().hex[:12]}",
                room_id=new_room.id,
                type="ELECTRICITY",
                serial_number=f"EM-{new_room.room_number}-{secrets.randbelow(9000)+1000}",
                initial_reading=0.0,
                current_reading=0.0,
                status="ACTIVE"
            )
            m_water = Meter(
                id=f"mtr_{uuid.uuid4().hex[:12]}",
                room_id=new_room.id,
                type="WATER",
                serial_number=f"WM-{new_room.room_number}-{secrets.randbelow(9000)+1000}",
                initial_reading=0.0,
                current_reading=0.0,
                status="ACTIVE"
            )
            db.add(m_elec)
            db.add(m_water)

        await db.commit()
    except Exception as e:
        await db.rollback()
        return error_response(code="IMPORT_EXECUTION_ERROR", message=f"Lỗi khi lưu dữ liệu: {str(e)}", status_code=400)

    return success_response({
        "message": "Import tòa nhà và danh sách phòng thành công!",
        "buildingsCreated": buildings_created,
        "floorsCreated": floors_created,
        "roomsCreated": rooms_created,
    }, status_code=201)


@router.get("/buildings/{building_id}", summary="Chi tiết tòa nhà theo ID (Public)")
async def get_building_by_id(building_id: str, db: AsyncSession = Depends(get_db)):
    await assert_public_building_access(db, building_id)
    stmt = (
        select(Building)
        .options(selectinload(Building.floors), selectinload(Building.rooms), selectinload(Building.configurations))
        .where(Building.id == building_id)
    )
    res = await db.execute(stmt)
    bld = res.scalar_one_or_none()
    if not bld:
        raise HTTPException(status_code=404, detail="Building not found")

    return success_response({
        "id": bld.id,
        "companyId": bld.company_id,
        "name": bld.name,
        "slug": bld.slug,
        "description": bld.description,
        "address": bld.address,
        "city": bld.city,
        "district": bld.district,
        "ward": bld.ward,
        "imageUrl": bld.image_url,
        "status": bld.status,
        "floors": [{"id": f.id, "floorNumber": f.floor_number, "name": f.name} for f in bld.floors],
        "rooms": [{"id": r.id, "roomNumber": r.room_number, "status": r.status, "baseRent": float(r.base_rent)} for r in bld.rooms],
        "configurations": [
            {
                "id": c.id,
                "version": c.version,
                "electricityUnitPrice": float(c.electricity_unit_price),
                "waterUnitPrice": float(c.water_unit_price),
                "internetPrice": float(c.internet_price),
            }
            for c in bld.configurations
        ],
    })


@router.post("/buildings", status_code=status.HTTP_201_CREATED, summary="Tạo mới tòa nhà")
async def create_building(
    payload: CreateBuildingPayload,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "SUPER_ADMIN")),
):
    comp_id = resolve_owner_company_id(current_user, payload.companyId)

    bld_slug = payload.slug
    if not bld_slug:
        import re
        bld_slug = re.sub(r'[^a-zA-Z0-9]+', '-', payload.name.lower()).strip('-') + f"-{uuid.uuid4().hex[:6]}"

    bld = Building(
        company_id=comp_id,
        name=payload.name,
        slug=bld_slug.lower(),
        description=payload.description,
        address=payload.address,
        city=payload.city or "Hồ Chí Minh",
        district=payload.district,
        ward=payload.ward,
        latitude=payload.latitude,
        longitude=payload.longitude,
        image_url=payload.imageUrl or "https://images.unsplash.com/photo-1545324418-cc1a3fa10c00?auto=format&fit=crop&w=1000&q=80",
        status="ACTIVE",
    )
    db.add(bld)
    await db.flush()

    # Create default configuration
    cfg = BuildingConfiguration(
        building_id=bld.id,
        version=1,
        effective_from=date.today(),
        electricity_unit_price=3500,
        water_unit_price=15000,
        internet_price=100000,
        garbage_price=50000,
        parking_fee_motorbike=100000,
        parking_fee_car=800000,
        cleaning_fee=150000,
        notes="Initial default property utility configuration",
    )
    db.add(cfg)
    await db.commit()
    await db.refresh(bld)

    return success_response({
        "id": bld.id,
        "name": bld.name,
        "slug": bld.slug,
        "address": bld.address,
        "status": bld.status,
    }, status_code=201)


@router.get("/buildings/{building_id}/configurations", summary="Danh sách cấu hình biểu phí dịch vụ tòa nhà (Private)")
async def get_building_configurations(
    building_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(get_current_user),
):
    await assert_private_building_access(db, current_user, building_id)
    res = await db.execute(
        select(BuildingConfiguration)
        .where(BuildingConfiguration.building_id == building_id)
        .order_by(BuildingConfiguration.version.desc())
    )
    cfgs = res.scalars().all()
    return success_response([
        {
            "id": c.id,
            "version": c.version,
            "effectiveFrom": c.effective_from.isoformat() if c.effective_from else None,
            "effectiveTo": c.effective_to.isoformat() if c.effective_to else None,
            "electricityUnitPrice": float(c.electricity_unit_price),
            "waterUnitPrice": float(c.water_unit_price),
            "internetPrice": float(c.internet_price),
            "garbagePrice": float(c.garbage_price),
            "parkingFeeMotorbike": float(c.parking_fee_motorbike),
            "parkingFeeCar": float(c.parking_fee_car),
            "cleaningFee": float(c.cleaning_fee),
            "notes": c.notes,
        }
        for c in cfgs
    ])


@router.post("/buildings/{building_id}/configurations", status_code=status.HTTP_201_CREATED)
async def create_building_configuration(
    building_id: str,
    payload: CreateBuildingConfigPayload,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "SUPER_ADMIN")),
):
    await assert_building_access(db, current_user, building_id)

    eff_str = payload.effectiveFrom or payload.effective_from or date.today().isoformat()
    try:
        eff_date = datetime.strptime(eff_str, "%Y-%m-%d").date()
    except Exception:
        eff_date = date.today()

    res = await db.execute(
        select(BuildingConfiguration)
        .where(BuildingConfiguration.building_id == building_id)
        .order_by(BuildingConfiguration.version.desc())
    )
    existing = res.scalars().all()
    next_version = len(existing) + 1

    if existing:
        latest_cfg = existing[0]
        if len(existing) == 1 and latest_cfg.notes == "Initial default property utility configuration":
            latest_cfg.effective_from = eff_date
            latest_cfg.electricity_unit_price = payload.electricityUnitPrice if payload.electricityUnitPrice is not None else (payload.electricity_unit_price or 3500)
            latest_cfg.water_unit_price = payload.waterUnitPrice if payload.waterUnitPrice is not None else (payload.water_unit_price or 15000)
            latest_cfg.notes = payload.notes or "Bản biểu giá gốc 2026"
            await db.commit()
            await db.refresh(latest_cfg)
            return success_response({"id": latest_cfg.id, "version": latest_cfg.version}, status_code=201)

        if eff_date <= latest_cfg.effective_from:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"INVALID_EFFECTIVE_DATE: New effective date {eff_date} must be strictly after previous version effective date {latest_cfg.effective_from}"
            )
        from datetime import timedelta
        latest_cfg.effective_to = eff_date - timedelta(days=1)

    cfg = BuildingConfiguration(
        building_id=building_id,
        version=next_version,
        effective_from=eff_date,
        electricity_unit_price=payload.electricityUnitPrice if payload.electricityUnitPrice is not None else (payload.electricity_unit_price or 3500),
        water_unit_price=payload.waterUnitPrice if payload.waterUnitPrice is not None else (payload.water_unit_price or 15000),
        internet_price=payload.internetPrice if payload.internetPrice is not None else (payload.internet_price or 100000),
        garbage_price=payload.garbagePrice if payload.garbagePrice is not None else (payload.garbage_price or 50000),
        parking_fee_motorbike=payload.parkingFeeMotorbike if payload.parkingFeeMotorbike is not None else (payload.parking_fee_motorbike or 100000),
        parking_fee_car=payload.parkingFeeCar if payload.parkingFeeCar is not None else (payload.parking_fee_car or 800000),
        cleaning_fee=payload.cleaningFee if payload.cleaningFee is not None else (payload.cleaning_fee or 150000),
        notes=payload.notes,
    )
    db.add(cfg)
    await db.commit()
    await db.refresh(cfg)

    return success_response({"id": cfg.id, "version": cfg.version}, status_code=201)


@router.patch("/buildings/{building_id}", summary="Cập nhật thông tin tòa nhà")
async def update_building(
    building_id: str,
    payload: UpdateBuildingPayload,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "SUPER_ADMIN")),
):
    bld = await assert_building_access(db, current_user, building_id)

    update_dict = payload.model_dump(exclude_unset=True)
    for k, v in update_dict.items():
        if k == "imageUrl":
            bld.image_url = v
        else:
            setattr(bld, k, v)

    await db.commit()
    await db.refresh(bld)

    return success_response({
        "id": bld.id,
        "name": bld.name,
        "slug": bld.slug,
        "description": bld.description,
        "address": bld.address,
        "district": bld.district,
        "city": bld.city,
        "status": bld.status,
    })


@router.delete("/buildings/{building_id}", summary="Xóa tòa nhà (chặn xóa nếu còn tầng/phòng/hợp đồng đang hoạt động)")
async def delete_building(
    building_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "SUPER_ADMIN")),
):
    bld = await assert_building_access(db, current_user, building_id)

    # Guard 1: check active contracts in any room of this building
    contract_stmt = (
        select(func.count(RentalContract.id))
        .join(Room, RentalContract.room_id == Room.id)
        .where(Room.building_id == building_id, RentalContract.status == "ACTIVE")
    )
    contract_res = await db.execute(contract_stmt)
    if (contract_res.scalar() or 0) > 0:
        raise HTTPException(status_code=400, detail="Không thể xóa tòa nhà đang có hợp đồng thuê hoạt động.")

    # Guard 2: check if building has floors or rooms
    floor_stmt = select(func.count(Floor.id)).where(Floor.building_id == building_id)
    floor_res = await db.execute(floor_stmt)
    if (floor_res.scalar() or 0) > 0:
        raise HTTPException(status_code=400, detail="Không thể xóa tòa nhà đang còn tầng. Vui lòng xóa các phòng và tầng trước.")

    room_stmt = select(func.count(Room.id)).where(Room.building_id == building_id)
    room_res = await db.execute(room_stmt)
    if (room_res.scalar() or 0) > 0:
        raise HTTPException(status_code=400, detail="Không thể xóa tòa nhà đang còn phòng. Vui lòng xóa các phòng trước.")

    await db.delete(bld)
    await db.commit()

    return success_response({"success": True, "message": "Đã xóa tòa nhà thành công"})


# ==========================================
# FLOORS
# ==========================================

@router.post("/buildings/{building_id}/floors", status_code=status.HTTP_201_CREATED, summary="Tạo mới tầng cho tòa nhà")
async def create_floor(
    building_id: str,
    payload: CreateFloorPayload,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "STAFF", "SUPER_ADMIN")),
):
    await assert_building_access(db, current_user, building_id)

    dup_stmt = select(Floor).where(Floor.building_id == building_id, Floor.floor_number == payload.floorNumber)
    dup_res = await db.execute(dup_stmt)
    if dup_res.scalar_one_or_none():
        raise HTTPException(status_code=400, detail=f"Tầng số {payload.floorNumber} đã tồn tại trong tòa nhà này")

    flr = Floor(
        building_id=building_id,
        floor_number=payload.floorNumber,
        name=payload.name,
        description=payload.description,
    )
    db.add(flr)
    await db.commit()
    await db.refresh(flr)

    return success_response({
        "id": flr.id,
        "buildingId": flr.building_id,
        "floorNumber": flr.floor_number,
        "name": flr.name,
        "description": flr.description,
    }, status_code=201)


@router.patch("/floors/{floor_id}", summary="Cập nhật thông tin tầng")
async def update_floor(
    floor_id: str,
    payload: UpdateFloorPayload,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "STAFF", "SUPER_ADMIN")),
):
    stmt = select(Floor).where(Floor.id == floor_id)
    res = await db.execute(stmt)
    flr = res.scalar_one_or_none()
    if not flr:
        raise HTTPException(status_code=404, detail="Không tìm thấy tầng")

    await assert_building_access(db, current_user, flr.building_id)

    if payload.floorNumber is not None and payload.floorNumber != flr.floor_number:
        dup_stmt = select(Floor).where(Floor.building_id == flr.building_id, Floor.floor_number == payload.floorNumber)
        dup_res = await db.execute(dup_stmt)
        if dup_res.scalar_one_or_none():
            raise HTTPException(status_code=400, detail=f"Tầng số {payload.floorNumber} đã tồn tại trong tòa nhà này")
        flr.floor_number = payload.floorNumber

    if payload.name is not None:
        flr.name = payload.name
    if payload.description is not None:
        flr.description = payload.description

    await db.commit()
    await db.refresh(flr)

    return success_response({
        "id": flr.id,
        "buildingId": flr.building_id,
        "floorNumber": flr.floor_number,
        "name": flr.name,
        "description": flr.description,
    })


@router.delete("/floors/{floor_id}", summary="Xóa tầng (chặn xóa nếu tầng còn phòng)")
async def delete_floor(
    floor_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "SUPER_ADMIN")),
):
    stmt = select(Floor).where(Floor.id == floor_id)
    res = await db.execute(stmt)
    flr = res.scalar_one_or_none()
    if not flr:
        raise HTTPException(status_code=404, detail="Không tìm thấy tầng")

    await assert_building_access(db, current_user, flr.building_id)

    # Guard: check if floor has rooms
    room_count_stmt = select(func.count(Room.id)).where(Room.floor_id == floor_id)
    room_count_res = await db.execute(room_count_stmt)
    room_count = room_count_res.scalar() or 0
    if room_count > 0:
        raise HTTPException(status_code=400, detail="Không thể xóa tầng đang chứa phòng. Vui lòng chuyển hoặc xóa các phòng trước.")

    await db.delete(flr)
    await db.commit()

    return success_response({"success": True, "message": "Đã xóa tầng thành công"})


# ==========================================
# ROOMS
# ==========================================

@router.get("/rooms", summary="Danh sách phòng (Public Explorer & Filter)")
async def list_rooms(
    buildingId: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    roomType: Optional[str] = Query(None),
    maxPrice: Optional[float] = Query(None),
    search: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    pageSize: int = Query(50, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
):
    offset = (page - 1) * pageSize
    rooms = await PropertyService.list_rooms(
        db,
        building_id=buildingId,
        status_filter=status,
        room_type=roomType,
        max_price=maxPrice,
        search=search,
        limit=pageSize,
        offset=offset,
    )
    return success_response(rooms)


@router.get("/rooms/slug/{slug}", summary="Chi tiết thông tin phòng theo slug")
async def get_room_by_slug(slug: str, db: AsyncSession = Depends(get_db)):
    stmt = (
        select(Room)
        .options(selectinload(Room.building), selectinload(Room.floor))
        .where(Room.slug == slug.lower())
    )
    res = await db.execute(stmt)
    room = res.scalar_one_or_none()
    if not room:
        raise HTTPException(status_code=404, detail="Room not found")

    return success_response({
        "id": room.id,
        "companyId": room.company_id,
        "buildingId": room.building_id,
        "floorId": room.floor_id,
        "roomNumber": room.room_number,
        "slug": room.slug,
        "roomType": room.room_type,
        "area": float(room.area),
        "baseRent": float(room.base_rent),
        "capacity": room.capacity,
        "status": room.status,
        "furnishing": room.furnishing,
        "description": room.description,
        "amenities": room.amenities or [],
        "images": room.images or [],
        "buildingName": room.building.name if room.building else None,
        "buildingAddress": room.building.address if room.building else None,
        "floorNumber": room.floor.floor_number if room.floor else None,
    })


@router.get("/rooms/{room_id}", summary="Chi tiết thông tin phòng")
async def get_room(room_id: str, db: AsyncSession = Depends(get_db)):
    room = await PropertyService.get_room_by_id(db, room_id)
    if not room:
        return error_response(code="ROOM_NOT_FOUND", message="Không tìm thấy phòng", status_code=404)
    return success_response(room)


@router.post("/rooms", status_code=status.HTTP_201_CREATED, summary="Tạo mới phòng")
async def create_room(
    payload: CreateRoomPayload,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "STAFF", "SUPER_ADMIN")),
):
    bld_id = payload.buildingId or payload.building_id
    if not bld_id:
        raise HTTPException(status_code=400, detail="MISSING_BUILDING_ID: buildingId is required")
    building = await assert_building_access(db, current_user, bld_id)

    r_num = payload.roomNumber or payload.room_number or "101"
    flr_id = payload.floorId or payload.floor_id
    r_type = payload.roomType or payload.room_type or "STUDIO"
    r_area = payload.area if payload.area is not None else (payload.area_sqm if payload.area_sqm is not None else 25.0)
    r_rent = payload.baseRent if payload.baseRent is not None else (payload.base_price_monthly if payload.base_price_monthly is not None else 5000000.0)
    if r_area <= 0:
        raise HTTPException(status_code=422, detail="Diện tích phòng phải lớn hơn 0")
    if r_rent < 0:
        raise HTTPException(status_code=422, detail="Giá thuê phòng không thể là số âm")
    r_cap = payload.max_occupants or payload.capacity or 2
    r_status = payload.status or "AVAILABLE"

    r_slug = payload.slug
    if not r_slug:
        r_slug = f"{building.slug}-{r_num.lower()}-{uuid.uuid4().hex[:6]}"

    room = Room(
        company_id=building.company_id,
        building_id=bld_id,
        floor_id=flr_id,
        room_number=r_num,
        slug=r_slug.lower(),
        room_type=r_type,
        area=r_area,
        base_rent=r_rent,
        capacity=r_cap,
        status=r_status,
        furnishing=payload.furnishing or "FULLY_FURNISHED",
        description=payload.description,
        amenities=payload.amenities or [],
        images=payload.images or [],
    )
    db.add(room)
    await db.flush()

    # Initialize default meters for room
    m_elec = Meter(
        id=f"mtr_{uuid.uuid4().hex[:12]}",
        room_id=room.id,
        type="ELECTRICITY",
        serial_number=f"EM-{room.room_number}-{secrets.randbelow(9000)+1000}",
        initial_reading=0.0,
        current_reading=0.0,
        status="ACTIVE"
    )
    m_water = Meter(
        id=f"mtr_{uuid.uuid4().hex[:12]}",
        room_id=room.id,
        type="WATER",
        serial_number=f"WM-{room.room_number}-{secrets.randbelow(9000)+1000}",
        initial_reading=0.0,
        current_reading=0.0,
        status="ACTIVE"
    )
    db.add(m_elec)
    db.add(m_water)
    await db.commit()
    await db.refresh(room)

    return success_response({
        "id": room.id,
        "roomNumber": room.room_number,
        "slug": room.slug,
        "status": room.status,
    }, status_code=201)


@router.patch("/rooms/{room_id}", summary="Cập nhật thông tin phòng")
async def update_room(
    room_id: str,
    payload: UpdateRoomPayload,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "STAFF", "SUPER_ADMIN")),
):
    room = await assert_room_access(db, current_user, room_id)

    update_dict = payload.model_dump(exclude_unset=True)
    for k, v in update_dict.items():
        if k in ("roomNumber", "room_number"):
            room.room_number = v
        elif k in ("roomType", "room_type"):
            room.room_type = v
        elif k in ("baseRent", "base_rent"):
            room.base_rent = v
        elif hasattr(room, k):
            setattr(room, k, v)

    await db.commit()
    await db.refresh(room)

    return success_response({
        "id": room.id,
        "roomNumber": room.room_number,
        "status": room.status,
        "baseRent": float(room.base_rent),
    })


@router.delete("/rooms/{room_id}", summary="Xóa phòng (chặn xóa nếu phòng đang có hợp đồng ACTIVE)")
async def delete_room(
    room_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: TokenPayload = Depends(require_role("OWNER", "SUPER_ADMIN")),
):
    room = await assert_room_access(db, current_user, room_id)

    # Guard: check if room has active contracts
    contract_stmt = select(func.count(RentalContract.id)).where(
        RentalContract.room_id == room_id,
        RentalContract.status == "ACTIVE"
    )
    contract_res = await db.execute(contract_stmt)
    if (contract_res.scalar() or 0) > 0:
        raise HTTPException(status_code=400, detail="Không thể xóa phòng đang có hợp đồng thuê hoạt động.")

    await db.delete(room)
    await db.commit()

    return success_response({"success": True, "message": "Đã xóa phòng thành công"})

